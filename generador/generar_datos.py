#!/usr/bin/env python3
"""Genera los archivos JSON usados por Consulta_Cecos.

Entradas esperadas:
  - Cecos_ASA.xlsx
  - Cecos_IA.xlsx
  - CECOS_Historico.xlsx (tambien acepta nombres con sufijos)
  - ASA_DataCostos.xlsx
  - IA_DataCostos.xlsx

Los importes no se publican. Los consumos se agrupan y se dividen en archivos
pequenos para que GitHub Pages cargue solo el CECO consultado.
"""

from __future__ import annotations

import argparse
import json
import math
import re
import shutil
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


ORIGINS = {"ASA": "ASA", "IA": "INTERANDINA"}
TARGET_CHUNK_BYTES = 550_000


def application_dir() -> Path:
    """Carpeta del script o del ejecutable portatil de Windows."""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def iso_date(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = clean(value)
    if not text:
        return ""
    match = re.match(r"(\d{4})-(\d{2})-(\d{2})", text)
    return match.group(0) if match else text


def number_or_none(value):
    if value in (None, ""):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if number == 0 or math.isnan(number):
        return None
    return int(number) if number.is_integer() else round(number, 6)


def compact_json(data) -> str:
    return json.dumps(data, ensure_ascii=False, separators=(",", ":"))


def write_json(path: Path, data) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(compact_json(data), encoding="utf-8")


def find_file(folder: Path, patterns: list[str]) -> Path:
    files = list(folder.glob("*.xlsx"))
    for pattern in patterns:
        exact = folder / pattern
        if exact.exists():
            return exact
    lowered = [(f, f.name.lower()) for f in files]
    for pattern in patterns:
        stem = Path(pattern).stem.lower()
        for file, name in lowered:
            if name.startswith(stem):
                return file
    raise FileNotFoundError("No se encontro: " + " / ".join(patterns))


def row_reader(path: Path, sheet_name: str | None = None, header_row: int = 1):
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    for _ in range(header_row - 1):
        next(rows, None)
    header = [clean(v).upper() for v in next(rows)]
    positions = {name: idx for idx, name in enumerate(header) if name}
    return workbook, worksheet, rows, positions


def preserve_restricted(script_dir: Path, input_dir: Path, explicit: str | None):
    candidates = []
    if explicit:
        candidates.append(Path(explicit))
    candidates.extend(
        [
            input_dir / "cecos.json",
            script_dir.parent / "cecos.json",
            script_dir.parent.parent / "cecos.json",
        ]
    )
    for candidate in candidates:
        if candidate.exists():
            try:
                raw = json.loads(candidate.read_text(encoding="utf-8"))
                return raw.get("uso_planillas", []), str(candidate)
            except Exception:
                continue
    return [], ""


def build_cecos(input_dir: Path, script_dir: Path, existing_cecos: str | None):
    records = []
    source_files = {}
    for sheet, origin in ORIGINS.items():
        path = find_file(input_dir, [f"Cecos_{sheet}.xlsx"])
        source_files[origin] = path.name
        wb, ws, rows, pos = row_reader(path, sheet)
        required = [
            "IDCONSUMIDOR",
            "JERARQUIA",
            "DESCRIPCION",
            "FECHA_INGRESO",
            "FECHA_BAJA",
            "AREA",
            "CUENTA_DESTINO",
        ]
        missing = [name for name in required if name not in pos]
        if missing:
            wb.close()
            raise ValueError(f"{path.name}: faltan columnas {', '.join(missing)}")
        for row in rows:
            code = clean(row[pos["IDCONSUMIDOR"]])
            if not code:
                continue
            records.append(
                {
                    "Origen": origin,
                    "Jerarquia": clean(row[pos["JERARQUIA"]]),
                    "Codigo": code,
                    "Descripcion": clean(row[pos["DESCRIPCION"]]),
                    "Fecha_inicio": iso_date(row[pos["FECHA_INGRESO"]]),
                    "Fecha_Fin": iso_date(row[pos["FECHA_BAJA"]]),
                    "Area": clean(row[pos["AREA"]]),
                    "Cuenta": clean(row[pos["CUENTA_DESTINO"]]),
                }
            )
        wb.close()
    restricted, preserved_from = preserve_restricted(script_dir, input_dir, existing_cecos)
    return {"cecos": records, "uso_planillas": restricted}, source_files, preserved_from


def build_history(input_dir: Path):
    path = find_file(input_dir, ["CECOS_Historico.xlsx", "CECOS_Historico(1).xlsx"])
    output = {
        "generado": date.today().isoformat(),
        "fuente": path.name,
        "ASA": {},
        "INTERANDINA": {},
        "rubros": {"ASA": {}, "INTERANDINA": {}},
        "rubros_filtro": {"ASA": {}, "INTERANDINA": {}},
    }
    rubro_codes = {"ASA": set(), "INTERANDINA": set()}
    warnings = []

    for sheet, origin in ORIGINS.items():
        wb, ws, rows, pos = row_reader(path, sheet)
        required = ["IDCCOSTO", "PARTIDA", "RUBRO", "FECHA"]
        missing = [name for name in required if name not in pos]
        if missing:
            wb.close()
            raise ValueError(f"{path.name}/{sheet}: faltan columnas {', '.join(missing)}")
        grouped = defaultdict(set)
        for row in rows:
            code = clean(row[pos["IDCCOSTO"]]).upper()
            if not code:
                continue
            partida = clean(row[pos["PARTIDA"]])
            rubro = clean(row[pos["RUBRO"]])
            movement_date = iso_date(row[pos["FECHA"]])
            grouped[code].add((partida, rubro, movement_date))
            if rubro:
                rubro_codes[origin].add(rubro.split()[0].upper())
        output[origin] = {
            code: [list(item) for item in sorted(values, key=lambda x: x[2], reverse=True)]
            for code, values in grouped.items()
        }
        wb.close()

    wb, ws, rows, pos = row_reader(path, "RUBRO", header_row=3)
    required = ["RUBRO", "FUNDO", "CULTIVO", "AREA", "FILTRO"]
    missing = [name for name in required if name not in pos]
    if missing:
        wb.close()
        raise ValueError(f"{path.name}/RUBRO: faltan columnas {', '.join(missing)}")
    mappings = {}
    for row in rows:
        rubro = clean(row[pos["RUBRO"]])
        if not rubro:
            continue
        code = rubro.split()[0].upper()
        mapping = [
            clean(row[pos["FUNDO"]]),
            clean(row[pos["CULTIVO"]]),
            clean(row[pos["AREA"]]),
            clean(row[pos["FILTRO"]]).upper(),
        ]
        if code in mappings and mappings[code] != mapping:
            warnings.append(f"Rubro {code}: clasificaciones diferentes; se uso la ultima")
        mappings[code] = mapping
    wb.close()

    for origin in ORIGINS.values():
        for code in sorted(rubro_codes[origin]):
            mapping = mappings.get(code)
            if not mapping:
                warnings.append(f"{origin}: rubro {code} sin clasificacion")
                continue
            fundo, cultivo, area, filtro = mapping
            output["rubros"][origin][code] = [area, fundo, cultivo]
            output["rubros_filtro"][origin][code] = [fundo, cultivo, area, filtro]
    return output, path.name, warnings


def build_consumption_source(path: Path, origin: str):
    wb, ws, rows, pos = row_reader(path)
    required = [
        "FECHA",
        "PERIODO",
        "IDCCOSTO",
        "CUENTA",
        "CANTIDAD",
        "DETALLE",
        "LABOR",
        "PARTIDA",
        "RUBRO",
    ]
    missing = [name for name in required if name not in pos]
    if missing:
        wb.close()
        raise ValueError(f"{path.name}: faltan columnas {', '.join(missing)}")

    data = {}
    periods = set()
    raw_rows = 0
    for row in rows:
        code = clean(row[pos["IDCCOSTO"]]).upper()
        if not code:
            continue
        raw_rows += 1
        period = clean(row[pos["PERIODO"]])
        periods.add(period)
        partida = clean(row[pos["PARTIDA"]])
        rubro = clean(row[pos["RUBRO"]])
        if code not in data:
            data[code] = {"summary": Counter(), "detail": {}, "count": 0}
        node = data[code]
        node["count"] += 1
        node["summary"][(partida, rubro, period)] += 1
        detail_key = (
            iso_date(row[pos["FECHA"]]),
            clean(row[pos["DETALLE"]]),
            clean(row[pos["CUENTA"]]),
            clean(row[pos["LABOR"]]),
            partida,
            rubro,
        )
        quantity = number_or_none(row[pos["CANTIDAD"]])
        if detail_key not in node["detail"]:
            node["detail"][detail_key] = [0, False, 0]
        aggregate = node["detail"][detail_key]
        if quantity is not None:
            aggregate[0] += quantity
            aggregate[1] = True
        aggregate[2] += 1
    sheet_name = ws.title
    wb.close()
    return data, periods, raw_rows, sheet_name


def finalize_consumptions(all_data, periods):
    period_index = {period: idx for idx, period in enumerate(periods)}
    final = {}
    for origin, data in all_data.items():
        final[origin] = {}
        for code, node in data.items():
            summary = defaultdict(lambda: [0] * len(periods))
            for (partida, rubro, period), count in node["summary"].items():
                if period in period_index:
                    summary[(partida, rubro)][period_index[period]] += count
            rows = [
                [partida, rubro, counts]
                for (partida, rubro), counts in sorted(summary.items(), key=lambda x: (x[0][1], x[0][0]))
            ]
            details = [
                [key[0], key[1], key[2], values[0] if values[1] else None, key[3], key[4], key[5], values[2]]
                for key, values in sorted(
                    node["detail"].items(),
                    key=lambda x: (x[0][0], x[0][4], x[0][2], x[0][1]),
                    reverse=True,
                )
            ]
            final[origin][code] = {"r": rows, "d": details, "_count": node["count"]}
    return final


def partition_consumptions(output_dir: Path, final, periods, sources, generated):
    folder = output_dir / "consumos"
    if folder.exists():
        shutil.rmtree(folder)
    folder.mkdir(parents=True, exist_ok=True)
    index = {
        "generado": generated,
        "fuente": " e ".join(sources),
        "periodos": periods,
        "ASA": {},
        "INTERANDINA": {},
    }
    files_created = {}
    for origin in ["ASA", "INTERANDINA"]:
        chunks = []
        current = {}
        current_size = 2
        for code in sorted(final[origin]):
            node = final[origin][code]
            movement_count = node.pop("_count")
            entry_size = len(compact_json({code: node}).encode("utf-8")) + 1
            if current and current_size + entry_size > TARGET_CHUNK_BYTES:
                chunks.append(current)
                current = {}
                current_size = 2
            chunk_number = len(chunks)
            current[code] = node
            current_size += entry_size
            index[origin][code] = [chunk_number, movement_count]
        if current:
            chunks.append(current)
        for number, chunk in enumerate(chunks):
            write_json(folder / f"{origin}_{number:03d}.json", chunk)
        files_created[origin] = len(chunks)
    write_json(folder / "index.json", index)
    return index, files_created


def parse_args():
    parser = argparse.ArgumentParser(description="Genera datos para Consulta_Cecos")
    script_dir = application_dir()
    parser.add_argument("--entradas", default=str(script_dir / "entradas"))
    parser.add_argument("--salida", default=str(script_dir / "salida"))
    parser.add_argument("--cecos-actual", default=None)
    return parser.parse_args()


def main():
    args = parse_args()
    script_dir = application_dir()
    input_dir = Path(args.entradas).resolve()
    output_dir = Path(args.salida).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    if not input_dir.exists():
        raise FileNotFoundError(f"No existe la carpeta de entradas: {input_dir}")

    print("1/4 Leyendo maestros de CECOs...")
    cecos, master_sources, preserved_from = build_cecos(input_dir, script_dir, args.cecos_actual)
    write_json(output_dir / "cecos.json", cecos)

    print("2/4 Generando historico...")
    history, history_source, warnings = build_history(input_dir)
    write_json(output_dir / "historico.json", history)

    print("3/4 Procesando movimientos ASA...")
    asa_file = find_file(input_dir, ["ASA_DataCostos.xlsx"])
    asa_data, asa_periods, asa_rows, asa_sheet = build_consumption_source(asa_file, "ASA")

    print("4/4 Procesando movimientos IA...")
    ia_file = find_file(input_dir, ["IA_DataCostos.xlsx"])
    ia_data, ia_periods, ia_rows, ia_sheet = build_consumption_source(ia_file, "INTERANDINA")
    periods = sorted((asa_periods | ia_periods) - {""})
    final = finalize_consumptions({"ASA": asa_data, "INTERANDINA": ia_data}, periods)
    generated = date.today().isoformat()
    source_labels = [f"{asa_file.name} (hoja {asa_sheet})", f"{ia_file.name} (hoja {ia_sheet})"]
    index, chunk_counts = partition_consumptions(output_dir, final, periods, source_labels, generated)

    report = {
        "generado": generated,
        "salida": str(output_dir),
        "maestros": {
            "registros": len(cecos["cecos"]),
            "uso_planillas": len(cecos["uso_planillas"]),
            "preservado_desde": preserved_from,
            "fuentes": master_sources,
        },
        "historico": {
            "fuente": history_source,
            "cecos_ASA": len(history["ASA"]),
            "cecos_INTERANDINA": len(history["INTERANDINA"]),
        },
        "consumos": {
            "periodos": periods,
            "filas_ASA": asa_rows,
            "filas_INTERANDINA": ia_rows,
            "cecos_ASA": len(index["ASA"]),
            "cecos_INTERANDINA": len(index["INTERANDINA"]),
            "archivos_ASA": chunk_counts["ASA"],
            "archivos_INTERANDINA": chunk_counts["INTERANDINA"],
        },
        "advertencias": warnings,
    }
    write_json(output_dir / "reporte_generacion.json", report)

    print("\nGENERACION TERMINADA")
    print(f"Carpeta: {output_dir}")
    print(f"CECOs: {len(cecos['cecos']):,}")
    print(f"Movimientos: {asa_rows + ia_rows:,}")
    print(f"Periodos: {', '.join(periods)}")
    print(f"Archivos de consumos: {chunk_counts['ASA'] + chunk_counts['INTERANDINA']}")
    if not preserved_from:
        print("ADVERTENCIA: no se encontro cecos.json anterior; uso_planillas quedo vacio.")
    for warning in warnings:
        print("ADVERTENCIA:", warning)


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print("\nERROR:", exc)
        sys.exit(1)
